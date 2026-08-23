# ============================================================================
# PROPIEDAD INTELECTUAL Y LICENCIA COMERCIAL CERRADA
# ============================================================================
# Autor Legal y Titular de Derechos: JAVIER ILLAN GONZALEZ
# Organización: ORANGE CREW
# Contacto: ILLANJAVIER9@GMAIL.COM
#
# ADVERTENCIA LEGAL (MÉXICO Y GLOBAL):
# Este código fuente y su arquitectura son propiedad intelectual exclusiva de
# JAVIER ILLAN GONZALEZ. Queda estrictamente prohibida su reproducción,
# distribución, modificación, ingeniería inversa, copia o uso comercial sin la
# autorización expresa y por escrito del autor. Obra protegida conforme a la
# Ley Federal del Derecho de Autor y tratados internacionales aplicables.
# ============================================================================
"""
reports/excel.py — Generación de archivos .xlsx con el formato de marca.

Segunda pieza movida fuera de app.py en el refactor por capas (ver
docs/PLAN_REFACTOR_CAPAS.md, fase 1). Reglas de esta capa:

- No sabe qué es "la empresa activa": para poner el nombre de la empresa
  como título de la hoja, recibe un `resolver_titulo` (una función sin
  argumentos que app.py registra una sola vez al arrancar, vía
  `set_resolver_titulo_empresa`). Así este módulo nunca importa nada de
  app.py ni de models/, y no hay riesgo de import circular.
- Recibe `logo_path` explícito (con default calculado aquí mismo a partir
  de la carpeta de este archivo) en vez de asumir una ruta fija.
"""
import io
import os
import sys

import pandas as pd

if getattr(sys, "frozen", False):
    _BASE_DIR = os.path.dirname(sys.executable)
else:
    _BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

LOGO_OLEASING_PNG_DEFAULT = os.path.join(_BASE_DIR, "assets", "o-leasing-logo.png")

_resolver_titulo_empresa = None


def set_resolver_titulo_empresa(fn) -> None:
    """app.py llama esto una sola vez al arrancar, pasándole una función
    que regrese el nombre de la empresa activa (por ejemplo:
    ``lambda: get_cfg('nombre_empresa', get_empresa_actual()['nombre'])``)."""
    global _resolver_titulo_empresa
    _resolver_titulo_empresa = fn


def formatear_hoja_excel(ws, df, currency_cols=None, pct_cols=None, int_cols=None,
                          freeze_header=True, titulo_empresa=None, logo_path=None):
    """Le da a una hoja de Excel recién escrita el mismo look del resto del
    sistema: el nombre de tu empresa como título (con 'Generado con
    O-Leasing' chiquito debajo, como crédito del programa, no al revés),
    encabezado en el verde de la marca, columnas con ancho ajustado al
    contenido, formato de moneda/porcentaje donde corresponde, encabezado
    congelado al hacer scroll, y filas alternadas para no perderse."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    logo_path = logo_path or LOGO_OLEASING_PNG_DEFAULT
    currency_cols = set(currency_cols or [])
    pct_cols      = set(pct_cols or [])
    int_cols      = set(int_cols or [])
    n_rows = len(df)
    n_cols = max(len(df.columns), 1)
    fila_header = 3 if titulo_empresa else 1  # fila donde vive el encabezado de columnas

    if titulo_empresa:
        ws.cell(row=1, column=1, value=titulo_empresa).font = Font(bold=True, size=13, color="1E5C4F", name="Calibri")
        ws.cell(row=2, column=1, value="Generado con O-Leasing · una app de Orange").font = Font(size=8, color="8A929C", italic=True, name="Calibri")
        if n_cols > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
        ws.row_dimensions[1].height = 22
        try:
            from openpyxl.drawing.image import Image as XLImage
            if os.path.exists(logo_path):
                _xl_logo = XLImage(logo_path)
                _xl_logo.height = 20
                _xl_logo.width = 20 * 2119 / 471
                _col_logo = get_column_letter(max(1, n_cols))
                _xl_logo.anchor = f"{_col_logo}1"
                ws.add_image(_xl_logo)
        except Exception:
            pass

    header_fill = PatternFill(start_color="1E5C4F", end_color="1E5C4F", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
    body_font   = Font(name="Calibri", size=10)
    alt_fill    = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    thin   = Side(style="thin", color="DCE0E5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, col_name in enumerate(df.columns, start=1):
        letra = get_column_letter(col_idx)
        cell = ws.cell(row=fila_header, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

        muestras = df[col_name].head(300)
        largo_valores = [len(f"{v:,.2f}") if isinstance(v, (int, float)) else len(str(v)) for v in muestras]
        max_len = max([len(str(col_name))] + largo_valores) if len(muestras) else len(str(col_name))
        ws.column_dimensions[letra].width = min(max(max_len + 2, 10), 42)

        num_fmt = None
        if col_name in currency_cols:
            num_fmt = '$#,##0.00;[RED]-$#,##0.00'
        elif col_name in pct_cols:
            num_fmt = '0.00"%"'
        elif col_name in int_cols:
            num_fmt = '#,##0'

        for row_idx in range(fila_header + 1, fila_header + n_rows + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.border = border
            c.font = body_font
            if num_fmt:
                c.number_format = num_fmt
            if (row_idx - fila_header) % 2 == 0:
                c.fill = alt_fill

    ws.row_dimensions[fila_header].height = 24
    if freeze_header:
        ws.freeze_panes = f"A{fila_header + 1}"
    ws.sheet_view.showGridLines = False
    if n_rows > 0:
        ws.auto_filter.ref = f"A{fila_header}:{get_column_letter(len(df.columns))}{fila_header + n_rows}"


def excel_con_formato(hojas: dict, currency_cols=None, pct_cols=None, int_cols=None,
                       incluir_titulo=True) -> io.BytesIO:
    """Genera un .xlsx con una o varias hojas, todas con el formato de
    'formatear_hoja_excel' y el nombre de tu empresa como título arriba de
    cada hoja. hojas: {'Nombre de hoja': dataframe, ...}. currency_cols/
    pct_cols/int_cols aplican a toda columna con ese nombre en cualquier
    hoja (pasa un dict {hoja: [...]} si necesitas distinguir).
    incluir_titulo=False para archivos que luego se vuelven a SUBIR al
    sistema (como la plantilla de carga masiva) — ahí la fila 1 tiene que
    quedar como encabezados de columna de verdad, no como título."""
    def _cols_para(hoja, spec):
        if spec is None:
            return []
        if isinstance(spec, dict):
            return spec.get(hoja, [])
        return spec

    _titulo_emp = None
    if incluir_titulo and _resolver_titulo_empresa:
        try:
            _titulo_emp = _resolver_titulo_empresa()
        except Exception:
            _titulo_emp = None

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        for nombre_hoja, df in hojas.items():
            hoja_segura = str(nombre_hoja)[:31]
            _startrow = 2 if _titulo_emp else 0
            df.to_excel(writer, sheet_name=hoja_segura, index=False, startrow=_startrow)
            ws = writer.sheets[hoja_segura]
            formatear_hoja_excel(
                ws, df,
                currency_cols=_cols_para(nombre_hoja, currency_cols),
                pct_cols=_cols_para(nombre_hoja, pct_cols),
                int_cols=_cols_para(nombre_hoja, int_cols),
                titulo_empresa=_titulo_emp,
            )
    buf.seek(0)
    return buf
